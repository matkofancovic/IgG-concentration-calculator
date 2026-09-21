#!/usr/bin/env python3
"""
GlycanAge plate run - GUI.

One window.  The Pippeting List goes in once and the NanoDrop export goes in
once; the concentration workbook and all three filled worksheets come out of
those two files.  Plates already seen are listed down the left, so a plate
part-way through is one click away from where it was left.

This module holds the core - reading the NanoDrop export and the plate
layout, and writing the concentration workbook.  The window itself is in
plate_run.py.

Requires: openpyxl, reportlab, pypdf   (pip install -r requirements.txt)
Run with: python igg_conc_gui.py
"""
__version__ = "2.2.0"

import os
import re
import sys
import traceback
from collections import OrderedDict

import tkinter as tk
from tkinter import ttk, filedialog, messagebox

try:
    import openpyxl
    from openpyxl.styles import PatternFill
except ImportError:
    sys.exit("openpyxl is not installed.  Run:  pip install openpyxl")

import ws_sheets
import ws_fill
import layout_colours


# ===========================================================================
#  Core
# ===========================================================================

ROWS = "ABCDEFGH"

# layout placeholders that occupy a well but are not a sample
FILLER = {"hidi", "hi-di", "formamide"}
WELL_RE = re.compile(r"[A-H](?:[1-9]|1[0-2])$")

# Well-type fills. Light tints so black text stays readable on screen and
# on a greyscale printout; keyed by the group name used in the layout.
FILLS = {
    "blank":    "FFC7CE",   # red
    "STAND_08": "C6EFCE",   # green
    "STAND_09": "BDD7EE",   # blue
    "STAND_10": "FFEB9C",   # yellow
}

# a blank must read below this, a standard within this range, or the
# validation step objects - these catch a wrong or stale plate layout
BLANK_MAX = 0.05
STAND_RANGE = (0.05, 1.0)


class BuildError(Exception):
    """Something is wrong with the inputs; message is shown to the user."""


def parse_num(s):
    """'0,2278' / '-8,307E-4' -> float, or None if not numeric."""
    s = (s or "").strip().replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def well_key(w):
    """Column-major sort: A1..H1, A2..H2, ..."""
    return (int(w[1:]), ROWS.index(w[0]))


def read_nanodrop(path):
    """-> (list of plate blocks, list of re-reads).

    Each block is an OrderedDict well -> reading.  A new block starts when
    the plate restarts at its first column; a well re-read part-way through
    a plate stays in the same block, and the later read wins.
    """
    try:
        with open(path, encoding="utf-8-sig", errors="replace") as fh:
            lines = [ln.rstrip("\r\n") for ln in fh if ln.strip()]
    except OSError as e:
        raise BuildError(f"Cannot read the NanoDrop file:\n{e}")
    if not lines:
        raise BuildError(f"{os.path.basename(path)} is empty.")

    header = [h.strip() for h in lines[0].split("\t")]
    idx = {name: i for i, name in enumerate(header)}
    need = ("Well", "conc.", "units", "Time")
    missing = [c for c in need if c not in idx]
    if missing:
        raise BuildError(
            f"{os.path.basename(path)} does not look like a NanoDrop plate export.\n"
            f"Missing column(s): {', '.join(missing)}")

    blocks, cur, rereads = [], OrderedDict(), []
    for ln in lines[1:]:
        f = ln.split("\t")
        if len(f) < len(header):
            f += [""] * (len(header) - len(f))
        well = f[idx["Well"]].strip()
        if not WELL_RE.match(well):
            continue
        rec = {
            "conc": parse_num(f[idx["conc."]]),
            "units": f[idx["units"]].strip() or "mg/ml",
            "time": f[idx["Time"]].strip(),
        }
        if well in cur:
            first_col = min(int(w[1:]) for w in cur)
            if int(well[1:]) <= first_col:          # plate starting over
                blocks.append(cur)
                cur = OrderedDict()
            else:                                    # re-read within this plate
                p = cur[well]
                rereads.append((len(blocks), well, p["time"], p["conc"],
                                rec["time"], rec["conc"]))
        cur[well] = rec
    if cur:
        blocks.append(cur)
    if not blocks:
        raise BuildError(f"No well readings found in {os.path.basename(path)}.")
    return blocks, rereads


def read_layout(path, sheet=None, keep_filler=False):
    """Read a plate-map sheet -> ({well: sample}, sheet name, conflicts).

    Handles both shapes in the pipeline: a single grid ('List 1') and a
    sheet repeating the same grid once per pipetting step ('Plate_layouts').

    keep_filler=True keeps HiDi and similar placeholders.  The worksheet
    sample list shows them; the concentration workbook ignores them.
    """
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
    except Exception as e:
        raise BuildError(f"Cannot open the plate layout:\n{e}")

    if sheet:
        if sheet not in wb.sheetnames:
            raise BuildError(f"No sheet named {sheet!r}.\n"
                             f"Sheets present: {', '.join(wb.sheetnames)}")
        name = sheet
    else:
        found = [s for s in ("Plate_layouts", "List 1") if s in wb.sheetnames]
        if not found:
            raise BuildError(
                f"{os.path.basename(path)} has no 'Plate_layouts' or 'List 1' sheet.\n"
                f"Sheets present: {', '.join(wb.sheetnames)}")
        name = found[0]

    ws = wb[name]
    layout, conflicts, cols = {}, [], None
    for row in ws.iter_rows():
        vals = [c.value for c in row]
        numbered = {j: int(str(v).strip()) for j, v in enumerate(vals)
                    if v is not None and str(v).strip().isdigit()}
        label = str(vals[0]).strip() if vals and vals[0] is not None else ""
        if len(numbered) >= 8 and label == "":
            cols = numbered                          # header row of a block
            continue
        if cols is None or label not in ROWS:
            continue
        for j, n in cols.items():
            if j >= len(vals) or vals[j] is None:
                continue
            val = str(vals[j]).strip()
            if not val or (val.lower() in FILLER and not keep_filler):
                continue
            well = f"{label}{n}"
            if well in layout and layout[well] != val:
                conflicts.append((well, layout[well], val))
            layout[well] = val
    if not layout:
        raise BuildError(f"Could not find a plate grid in sheet {name!r}.")
    return layout, name, conflicts


def classify(sample):
    """-> ('standard', group) | ('blank', None) | ('sample', None)"""
    m = re.match(r"(STAND_\d+)_\d+$", sample)
    if m:
        return "standard", m.group(1)
    if sample.lower().startswith("blank"):
        return "blank", None
    return "sample", None


def pick_block(blocks, layout):
    """Choose the plate whose wells best match the layout."""
    keys = set(layout)
    return max((len(set(b) & keys), -abs(len(b) - len(keys)), i)
               for i, b in enumerate(blocks))[2]


def compact(rows):
    """[2,3,4,7] -> 'C2:C4,C7'"""
    out, i = [], 0
    while i < len(rows):
        j = i
        while j + 1 < len(rows) and rows[j + 1] == rows[j] + 1:
            j += 1
        out.append(f"C{rows[i]}" if i == j else f"C{rows[i]}:C{rows[j]}")
        i = j + 1
    return ",".join(out)


def mean_concentrations(nanodrop_path, layout, plate=None):
    """-> (dbs_mean, standards_mean) in mg/ml, from a NanoDrop export.

    DBS is the mean of the sample wells; standards is the mean of every
    standard well pooled.  Blanks and unmapped wells are excluded from both.
    """
    blocks, _ = read_nanodrop(nanodrop_path)
    wells = blocks[(plate - 1) if plate else pick_block(blocks, layout)]
    dbs, std = [], []
    for w, rec in wells.items():
        name = layout.get(w)
        # layout may have been read with keep_filler=True for the worksheet
        # grid; HiDi wells hold no IgG and must not enter either mean
        if not name or rec["conc"] is None or name.lower() in FILLER:
            continue
        kind, _ = classify(name)
        if kind == "sample":
            dbs.append(rec["conc"])
        elif kind == "standard":
            std.append(rec["conc"])
    if not dbs and not std:
        raise BuildError("No sample or standard wells matched the layout - "
                         "is this the right NanoDrop file for this plate?")
    return (sum(dbs) / len(dbs) if dbs else None,
            sum(std) / len(std) if std else None)


def validate(layout, wells, stand_groups):
    """-> list of complaints.  Mostly catches a wrong or stale plate layout."""
    problems = []
    for g, members in sorted(stand_groups.items()):
        if len(members) != 4:
            problems.append(f"{g} has {len(members)} replicate(s), expected 4")
    if not stand_groups:
        problems.append("No standards found in the layout (expected STAND_xx_n names)")

    blanks = [w for w, s in layout.items() if classify(s)[0] == "blank"]
    if not blanks:
        problems.append("No blanks found in the layout")
    for w in sorted(blanks, key=well_key):
        c = wells.get(w, {}).get("conc")
        if c is not None and c > BLANK_MAX:
            problems.append(
                f"{w} is {layout[w]} in the layout but reads {c:.4g} mg/ml "
                f"- too high for a blank (wrong plate layout?)")

    lo, hi = STAND_RANGE
    for g, members in sorted(stand_groups.items()):
        for w in members:
            c = wells.get(w, {}).get("conc")
            if c is not None and not (lo <= c <= hi):
                problems.append(
                    f"{w} is {layout[w]} but reads {c:.4g} mg/ml "
                    f"- outside the expected {lo}-{hi} mg/ml")
    return problems


def build(nanodrop_path, layout_path, out_path, sheet=None,
          plate=None, sample_label="DBS", strict=True, log=print):
    """Do the whole job.  `log` receives one report line at a time."""
    blocks, rereads = read_nanodrop(nanodrop_path)
    layout, sheet_used, conflicts = read_layout(layout_path, sheet)

    chosen = (plate - 1) if plate else pick_block(blocks, layout)
    if not 0 <= chosen < len(blocks):
        raise BuildError(f"Plate {plate} out of range: the file holds {len(blocks)}.")
    wells = blocks[chosen]

    log(f"Version  : {__version__}")
    log(f"NanoDrop : {os.path.basename(nanodrop_path)}")
    log(f"Layout   : {os.path.basename(layout_path)}  [{sheet_used}]")
    if len(blocks) > 1:
        log(f"Plate    : block {chosen + 1} of {len(blocks)} in the file "
            f"({len(wells)} wells)")

    # ---- work out what goes where before writing anything ----------------
    order = sorted({w for w in set(wells) | set(layout) if WELL_RE.match(w)},
                   key=well_key)
    stand_groups = {}
    for w in order:
        s = layout.get(w)
        if s and w in wells:
            kind, grp = classify(s)
            if kind == "standard":
                stand_groups.setdefault(grp, []).append(w)

    problems = validate(layout, wells, stand_groups)
    for p in problems:
        log(f"  !! {p}")
    if problems and strict:
        raise BuildError(
            "The layout does not match the measurements:\n\n  - "
            + "\n  - ".join(problems)
            + "\n\nThis usually means the wrong plate layout file was picked.\n"
              "Untick 'Stop if the layout does not match' to write it anyway.")

    # ---- write -----------------------------------------------------------
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Well ", "Sample ID", "conc. ", "units "])
    ws["C1"].number_format = "0.0000"

    def fill_row(row, key, first=1, last=4):
        rgb = FILLS.get(key)
        if not rgb:
            return
        pf = PatternFill("solid", fgColor=rgb)
        for col in range(first, last + 1):
            ws.cell(row=row, column=col).fill = pf

    stand_rows, sample_rows, unmeasured, unmapped = {}, [], [], []
    for w in order:
        rec, sample = wells.get(w), layout.get(w)
        if rec is None:
            unmeasured.append(w)
            continue
        if sample is None:
            unmapped.append(w)
        ws.append([w, sample, rec["conc"], rec["units"]])
        r = ws.max_row
        ws.cell(row=r, column=3).number_format = "0.0000"
        if not sample:
            continue
        kind, grp = classify(sample)
        if kind == "standard":
            stand_rows.setdefault(grp, []).append(r)
            fill_row(r, grp)
        elif kind == "blank":
            fill_row(r, "blank")
        else:
            sample_rows.append(r)

    if not sample_rows:
        raise BuildError("No sample wells found - check the plate layout.")

    start = ws.max_row + 2
    ws.cell(row=start, column=2, value=sample_label)
    ws.cell(row=start, column=3,
            value=f"=AVERAGE({compact(sample_rows)})").number_format = "0.0000"
    ws.cell(row=start, column=4, value="mg/ml")
    for k, name in enumerate(sorted(stand_rows), 1):
        ws.cell(row=start + k, column=2, value=name)
        ws.cell(row=start + k, column=3,
                value=f"=AVERAGE({compact(sorted(stand_rows[name]))})"
                ).number_format = "0.0000"
        ws.cell(row=start + k, column=4, value="mg/ml")
        fill_row(start + k, name, first=2)

    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 8.88671875

    try:
        wb.save(out_path)
    except PermissionError:
        raise BuildError(
            f"Cannot write {os.path.basename(out_path)}.\n\n"
            "It is probably open in Excel - close it and try again.")
    except OSError as e:
        raise BuildError(f"Cannot write the output file:\n{e}")

    # ---- report ----------------------------------------------------------
    log("")
    log(f"Wells written : {len(order) - len(unmeasured)}")
    log(f"Samples       : {len(sample_rows)}")
    log("Standards     : " + ", ".join(f"{k} ({len(v)})"
                                       for k, v in sorted(stand_rows.items())))
    blanks = sorted((w for w, s in layout.items() if classify(s)[0] == "blank"),
                    key=well_key)
    if blanks:
        log("Blanks        : " + ", ".join(
            f"{w} {layout[w]} = {wells[w]['conc']:.4g}"
            for w in blanks if w in wells))
    for b, w, t1, c1, t2, c2 in rereads:
        if b == chosen:
            log(f"  re-read {w}: {t1} {c1:.4g}  ->  {t2} {c2:.4g}   (later read kept)")
    for w, v1, v2 in conflicts[:10]:
        log(f"  !! layout conflict {w}: {v1!r} vs {v2!r}")
    if unmeasured:
        log("  in layout, not measured: " + ", ".join(unmeasured))
    if unmapped:
        log("  measured, not in layout: " + ", ".join(unmapped)
            + "   (written with no Sample ID, excluded from averages)")
    low = [(w, wells[w]["conc"]) for w in order
           if w in wells and layout.get(w) and classify(layout[w])[0] == "sample"
           and wells[w]["conc"] is not None and wells[w]["conc"] < BLANK_MAX]
    if low:
        log(f"  low (<{BLANK_MAX} mg/ml): "
            + ", ".join(f"{w} {c:.4g}" for w, c in low))
    log("")
    log(f"Saved: {out_path}")
    return out_path


# ===========================================================================
#  GUI
# ===========================================================================

def norm(path):
    """Windows-native form of a path.

    Tk's file dialogs hand back forward slashes, so a share comes through as
    //host/share/file.  Python opens that happily but os.startfile and other
    Windows APIs do not, so every path from a dialog goes through here.
    """
    return os.path.normpath(path) if path else path


class App(ttk.Frame):
    def __init__(self, master):
        super().__init__(master, padding=10)
        self.grid(sticky="nsew")
        master.columnconfigure(0, weight=1)
        master.rowconfigure(0, weight=1)
        self.columnconfigure(0, weight=1)
        self.rowconfigure(0, weight=9)      # the steps get the room
        self.rowconfigure(2, weight=1)      # the report pane

        import plate_run
        self.view = plate_run.MainView(self, self)
        self.view.grid(row=0, column=0, sticky="nsew")
        self.panel = self.view.panel

        ttk.Label(self, text="Report", style="Muted.TLabel").grid(
            row=1, column=0, sticky="w", pady=(10, 4))
        wrap = ttk.Frame(self)
        wrap.grid(row=2, column=0, sticky="nsew")
        wrap.columnconfigure(0, weight=1)
        wrap.rowconfigure(0, weight=1)
        self.log = tk.Text(wrap, height=4, wrap="none", state="disabled",
                           font=("Consolas", 10), relief="flat", borderwidth=0,
                           highlightthickness=1, highlightbackground="#DCE5E0",
                           highlightcolor="#DCE5E0", background="#FFFFFF",
                           foreground="#09341F", padx=12, pady=10)
        self.log.grid(row=0, column=0, sticky="nsew")
        sb = ttk.Scrollbar(wrap, orient="vertical", command=self.log.yview)
        sb.grid(row=0, column=1, sticky="ns")
        self.log.configure(yscrollcommand=sb.set)
        self.log.tag_configure("warn", foreground="#b00000")
        self.configure(padding=0)

        self.status = ttk.Label(self, text="Open a plate to start.",
                                anchor="w", style="Muted.TLabel")
        self.status.grid(row=3, column=0, sticky="ew", pady=(8, 0))

    # -- shared helpers ----------------------------------------------------

    def say(self, line=""):
        tag = "warn" if line.strip().startswith("!!") else ""
        self.log.configure(state="normal")
        self.log.insert("end", line + "\n", tag)
        self.log.see("end")
        self.log.configure(state="disabled")
        self.update_idletasks()

    def clear(self):
        self.log.configure(state="normal")
        self.log.delete("1.0", "end")
        self.log.configure(state="disabled")

    def confirm_overwrite(self, path):
        return not os.path.exists(path) or messagebox.askyesno(
            "Overwrite?", f"{os.path.basename(path)} already exists.\n\nOverwrite it?")

    def go(self, btn, out, open_after, work):
        """Run `work`, funnelling errors to the report pane and a dialog."""
        self.clear()
        btn.state(["disabled"])
        self.status.config(text="Working...")
        try:
            work()
        except (BuildError, ws_fill.FillError) as e:
            self.say("")
            self.say("!! " + str(e).replace("\n", "\n   "))
            self.status.config(text="Not written - see the report.")
            messagebox.showerror("Could not build it", str(e))
        except Exception:
            self.say("")
            self.say("!! Unexpected error:")
            self.say(traceback.format_exc())
            self.status.config(text="Failed - see the report.")
        else:
            self.status.config(text=f"Done - {os.path.basename(out)}")
            if open_after:
                try:
                    os.startfile(norm(out))
                except Exception as e:
                    self.say(f"  (could not open the file: {e})")
        finally:
            btn.state(["!disabled"])


def main():
    root = tk.Tk()
    root.title(f"GlycanAge - plate run  v{__version__}")
    root.minsize(1320, 920)          # the plate list sits beside the form
    import plate_run
    plate_run.apply_theme(root)
    App(root)
    root.mainloop()


USAGE = """Batch use:

  workbook:  <nanodrop.txt> <layout.xlsx> <output.xlsx> [--lenient]
  fill:      --fill <layout.xlsx> <out folder> [--nanodrop <file.txt>]
                     [--initials MF] [--blanks <folder>] [--aliquot 40]
                     [--pages isolation,deglyco,cleanup]
  sheets:    --sheets <layout.xlsx> <output.pdf> [--nanodrop <file.txt>]
                      [--aliquot 40] [--pages isolation,deglyco,cleanup]
                      [--initials MF]

With no arguments the window opens instead."""


def _opt(argv, name, default=None):
    """Value of '--name x', or default."""
    return argv[argv.index(name) + 1] if name in argv and \
        argv.index(name) + 1 < len(argv) else default


def cli(argv):
    """Batch mode - see USAGE."""
    try:
        if "--fill" in argv:
            consumed = {"--nanodrop", "--aliquot", "--pages", "--initials",
                        "--blanks"}
            skip = {argv[argv.index(o) + 1] for o in consumed
                    if o in argv and argv.index(o) + 1 < len(argv)}
            args = [a for a in argv if not a.startswith("--") and a not in skip]
            if len(args) != 2:
                print(USAGE)
                return 2
            layout_path, out_dir = args
            layout, sheet_used, _ = read_layout(layout_path, keep_filler=True)
            nd = _opt(argv, "--nanodrop")
            pages = _opt(argv, "--pages")
            ws_fill.fill_all(
                ws_fill.discover(_opt(argv, "--blanks")), out_dir, layout,
                batch=ws_sheets.batch_from_path(layout_path),
                numbers={}, initials=(_opt(argv, "--initials") or "").upper(),
                avg_conc=mean_concentrations(nd, layout) if nd else None,
                aliquot_ul=float(_opt(argv, "--aliquot", 40)),
                colours=layout_colours.read_colours(layout_path, sheet_used),
                pages=[x.strip() for x in pages.split(",")] if pages else None)
            return 0

        if "--sheets" in argv:
            nd = _opt(argv, "--nanodrop")
            consumed = {"--nanodrop", "--aliquot", "--pages", "--initials"}
            skip = {argv[argv.index(o) + 1] for o in consumed
                    if o in argv and argv.index(o) + 1 < len(argv)}
            args = [a for a in argv
                    if not a.startswith("--") and a not in skip]
            if len(args) != 2:
                print(USAGE)
                return 2
            layout_path, out = args
            pages = (_opt(argv, "--pages") or "").split(",") if \
                _opt(argv, "--pages") else None
            layout, sheet_used, _ = read_layout(layout_path, keep_filler=True)
            avg = mean_concentrations(nd, layout) if nd else None
            ws_sheets.build_worksheet_pack(
                layout, out,
                batch=ws_sheets.batch_from_path(layout_path),
                source_name=os.path.basename(layout_path),
                sheet_name=sheet_used, avg_conc=avg,
                aliquot_ul=float(_opt(argv, "--aliquot", 40)),
                initials=(_opt(argv, "--initials") or "").upper(),
                pages=[p.strip() for p in pages] if pages else None)
            return 0

        args = [a for a in argv if not a.startswith("--")]
        if len(args) != 3:
            print(__doc__)
            print(USAGE)
            return 2
        build(*args, strict="--lenient" not in argv)
    except (BuildError, ws_fill.FillError) as e:
        print(f"\nERROR: {e}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    # No arguments -> GUI.  Arguments -> batch, for scripting a whole run.
    sys.exit(cli(sys.argv[1:]) if len(sys.argv) > 1 else main())
